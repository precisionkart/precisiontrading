"""output.py — CSV / XLSX / terminal output (spec Section 7 + product brief).

  * write_csv      — one UTF-8 CSV per list, full precision.
  * write_xlsx     — one workbook, three sheets (Focus/Targets/Earnings), each
                     with tabular number formatting, a green gradient on the
                     Score column, and a frozen header row.
  * print_table    — rich terminal table with a colour-coded Score column.
"""

from __future__ import annotations

import logging
import os
from typing import Optional

import pandas as pd

logger = logging.getLogger("pinpoint.output")


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def write_csv(df: pd.DataFrame, path: str) -> Optional[str]:
    if df is None or len(df) == 0:
        return None
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8")
    logger.info("wrote %s (%d rows)", path, len(df))
    return path


def write_all_csv(lists: dict[str, pd.DataFrame], out_dir: str, date: str) -> dict[str, str]:
    paths = {}
    for name, df in lists.items():
        p = write_csv(df, os.path.join(out_dir, f"{name}_{date}.csv"))
        if p:
            paths[name] = p
    return paths


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
def write_xlsx(focus: pd.DataFrame, targets: pd.DataFrame, earnings: pd.DataFrame,
               path: str) -> Optional[str]:
    """Single workbook, three formatted sheets. Returns the path (or None)."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheets = {"Focus": focus, "Targets": targets, "Earnings": earnings}
    if all(df is None or len(df) == 0 for df in sheets.values()):
        return None

    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="111827")
    num_align = Alignment(horizontal="right")

    for sheet_name, df in sheets.items():
        ws = wb.create_sheet(sheet_name)
        if df is None or len(df) == 0:
            ws["A1"] = "(empty)"
            continue
        cols = list(df.columns)
        ws.append(cols)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for _, row in df.iterrows():
            ws.append([_xl_value(v) for v in row.tolist()])

        # formatting: tabular numbers, widths
        for ci, col in enumerate(cols, start=1):
            letter = get_column_letter(ci)
            ws.column_dimensions[letter].width = max(10, min(28, len(str(col)) + 6))
            for cell in ws[letter][1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
                    cell.alignment = num_align

        # green gradient on the Score column
        if "pinpoint_score" in cols:
            sidx = cols.index("pinpoint_score") + 1
            letter = get_column_letter(sidx)
            rng = f"{letter}2:{letter}{ws.max_row}"
            ws.conditional_formatting.add(rng, ColorScaleRule(
                start_type="min", start_color="FFFFFF",
                end_type="max", end_color="16A34A"))

        ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)
    logger.info("wrote %s", path)
    return path


def _xl_value(v):
    if v is None:
        return ""
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return str(v)
    try:
        import math
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return ""
    except Exception:
        pass
    try:
        import numpy as np
        if isinstance(v, (np.integer,)):
            return int(v)
        if isinstance(v, (np.floating,)):
            return float(v)
        if isinstance(v, (np.bool_,)):
            return bool(v)
    except Exception:
        pass
    return v


# ---------------------------------------------------------------------------
# Terminal (rich)
# ---------------------------------------------------------------------------
def print_table(title: str, df: pd.DataFrame, cols: Optional[list[str]] = None,
                score_col: str = "pinpoint_score") -> None:
    print(f"\n=== {title} ({0 if df is None else len(df)}) ===")
    if df is None or len(df) == 0:
        print("  (empty)")
        return
    view = df[cols] if cols else df
    try:
        from rich.console import Console
        from rich.table import Table

        console = Console()
        table = Table(show_header=True, header_style="bold", box=None, pad_edge=False)
        for c in view.columns:
            justify = "right" if c != view.columns[0] else "left"
            table.add_column(str(c), justify=justify)
        for _, r in view.iterrows():
            cells = []
            for c in view.columns:
                cells.append(_rich_cell(c, r[c], score_col))
            table.add_row(*cells)
        console.print(table)
    except Exception:  # noqa: BLE001
        with pd.option_context("display.max_columns", None, "display.width", 180):
            print(view.to_string(index=False))


def _rich_cell(col: str, val, score_col: str) -> str:
    if isinstance(val, float):
        txt = "-" if val != val else f"{val:,.2f}"
    else:
        txt = str(val)
    if col == score_col and isinstance(val, (int, float)) and val == val:
        color = "green3" if val >= 8 else "yellow3" if val >= 5 else "grey58"
        return f"[{color}]{txt}[/{color}]"
    return txt
