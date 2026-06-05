"""Tests for chart/report/output rendering — all offline (fixtures, no network)."""

import os

import pandas as pd
import pytest

from pinpoint import sample_data
from pinpoint import charts, report, output
from pinpoint import regime as regime_mod


def _focus_df():
    return pd.DataFrame([{
        "ticker": "TST", "company": "Test Co", "sector": "Technology",
        "price": 70.5, "rs": 95.0, "pattern": "High & tight flag", "pattern_bars": 15,
        "entry_trigger": 70.5, "stop": 67.19, "stop_kind": "quarter", "risk": 3.31,
        "measured_target": 106.6, "reward_risk": 10.9, "continuity": 1.0,
        "beach_ball": False, "pinpoint_score": 11.5, "n_layers": 9,
        "layers": "Valid bullish pattern (+1.5) | R:R >= 5:1 (+1.0)",
    }])


def test_chart_renders_png(tmp_path):
    daily = sample_data.ohlcv_fixture("TST")
    ann = charts.ChartAnnotation(entry=70.5, stop=67.19, target=106.6, reward_risk=10.9,
                                 pattern_label="High & tight flag", pattern_bars=15,
                                 score=11.5, rs=95)
    out = str(tmp_path / "tst.png")
    path = charts.render_focus_chart("TST", daily, out, ann)
    assert path is not None and os.path.exists(path)
    assert os.path.getsize(path) > 5000


def test_report_renders_html(tmp_path):
    reg = regime_mod.from_fundaments(sample_data.regime_fixture())
    focus = _focus_df()
    targets = pd.DataFrame([{"ticker": "TST", "sector": "Technology", "price": 70.5,
                             "rs": 95.0, "stage": "Stage 2 (advancing)", "growth": "EPS yr 40%",
                             "n_layers": 9, "pinpoint_score": 11.5}])
    earnings = pd.DataFrame([{"ticker": "ERN", "sector": "Tech", "price": 50.0,
                              "gap": 8.0, "rel_volume": 3.0, "rs": 80.0}])
    out = str(tmp_path / "report.html")
    path = report.render_report(focus, targets, earnings, reg, out, "10:00", "2026-06-06")
    assert os.path.exists(path)
    html = open(path, encoding="utf-8").read()
    assert "Pinpoint" in html and "TST" in html
    assert "316" not in html        # sanity: only our data
    assert "R:R" in html


def test_csv_and_xlsx(tmp_path):
    focus = _focus_df()
    csv_path = output.write_csv(focus, str(tmp_path / "focus.csv"))
    assert csv_path and os.path.exists(csv_path)

    xlsx_path = output.write_xlsx(focus, focus.rename(columns={}), focus,
                                  str(tmp_path / "wb.xlsx"))
    assert xlsx_path and os.path.exists(xlsx_path)
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    assert set(wb.sheetnames) == {"Focus", "Targets", "Earnings"}


def test_write_csv_empty_returns_none(tmp_path):
    assert output.write_csv(pd.DataFrame(), str(tmp_path / "x.csv")) is None
